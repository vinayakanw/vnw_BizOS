from pathlib import Path


class ParseError(Exception):
    pass


def _read_text(path: Path) -> str:
    return path.read_bytes().decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        raise ParseError("No extractable text (scanned PDF? OCR is not supported yet)")
    return text


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text(path: Path, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in {".txt", ".md", ".csv"}:
        text = _read_text(path)
    elif ext == ".pdf":
        text = _read_pdf(path)
    elif ext == ".docx":
        text = _read_docx(path)
    else:
        raise ParseError(f"Unsupported file type: {ext}")
    if not text.strip():
        raise ParseError("File contains no text")
    return text
